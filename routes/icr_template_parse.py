"""
ICR template parser: POST /parse-template (called by reporting_system_node
/api/icr/admin/parse-template with the user's token). Ported from the main branch
(commit 7af2241b "word parse": the /parse-template alias of /word-template/analyze); only the
parsing code, none of the other word-template routes. The helpers below are unchanged.

Only reporting Node may call it: Node runs the per-user upload quota (R-WAPT02) before
forwarding, so a direct call from a browser would skip that count. Node adds X-Internal-Token,
a 60-second HS256 token signed with the shared JWT_SECRET for this route and this user; users
cannot make one. The user's own JWT and the ICR Templates page check still apply.

Settings (.env), read per request:
  UPLOAD_MAX_FILE_MB             largest template accepted            (default 50, as reporting Node)
  WORD_TEMPLATE_MAX_UNZIPPED_MB  largest unzipped size of a .docx     (default 200, guards zip bombs)
"""
from fastapi import APIRouter, UploadFile, File, HTTPException, Request
from fastapi.responses import JSONResponse
from io import BytesIO
import logging
import os
import re
import zipfile
from docx import Document
from docx.document import Document as DocType
from docx.oxml.text.paragraph import CT_P
from docx.oxml.table import CT_Tbl
from docx.table import Table
from docx.text.paragraph import Paragraph
import jwt as pyjwt
from utils.auth import JWT_ALGORITHM, JWT_CLOCK_LEEWAY_SECONDS, JWT_SECRET, resolve_jwt_token_and_payload

router = APIRouter(tags=["icr-template"])
logger = logging.getLogger(__name__)

_MB = 1024 * 1024


def _mb_setting(name: str, default_mb: float) -> int:
	try:
		mb = float(os.getenv(name, ""))
	except ValueError:
		mb = 0
	return int((mb if mb > 0 else default_mb) * _MB)


INTERNAL_AUDIENCE = "reporting-python:parse-template"


def _sent_by_reporting_node(request: Request) -> bool:
	"""X-Internal-Token is valid, unexpired, for this route and for the logged-in user."""
	raw = (request.headers.get("x-internal-token") or "").strip()
	if not raw:
		return False
	try:
		claims = pyjwt.decode(
			raw, JWT_SECRET, algorithms=[JWT_ALGORITHM], audience=INTERNAL_AUDIENCE,
			options={"require": ["exp", "aud", "sub"]}, leeway=JWT_CLOCK_LEEWAY_SECONDS,
		)
	except pyjwt.PyJWTError:
		return False
	# login tokens carry no aud / exp of this kind; checked again explicitly
	if claims.get("aud") != INTERNAL_AUDIENCE or "exp" not in claims:
		return False
	_, user = resolve_jwt_token_and_payload(request)
	user_id = (user or {}).get("id") or (user or {}).get("sub")
	return bool(user_id) and str(claims.get("sub", "")).strip().lower() == str(user_id).strip().lower()


def _bad_request(message: str) -> JSONResponse:
	return JSONResponse(status_code=400, content={'success': False, 'error': message})


def extract_all_placeholders_from_docx_bytes(file_content: bytes):
	"""Extract all {{ ... }} placeholders from all XML parts inside DOCX."""
	try:
		all_text = []
		with zipfile.ZipFile(BytesIO(file_content)) as zf:
			for name in zf.namelist():
				if name.endswith('.xml'):
					try:
						all_text.append(zf.read(name).decode('utf-8', errors='ignore'))
					except Exception:
						continue
		combined = '\n'.join(all_text)
		matches = re.findall(r'\{\{\s*[^{}]+?\s*\}\}', combined)
		# Keep insertion order and uniqueness
		unique = list(dict.fromkeys(matches))
		return unique, len(matches)
	except Exception as e:
		logger.warning(f"Failed to extract raw placeholders from DOCX zip: {e}")
		return [], 0


def extract_sections_from_docx(doc: DocType):
	"""Extract sections (headings, paragraphs, tables) from Word document with formatting"""
	sections = []
	section_id = 0
	
	current_section = None
	
	# Store original document reference for later use
	template_doc = doc
	
	for element in doc.element.body:
		if isinstance(element, CT_P):
			paragraph = Paragraph(element, doc)
			text = paragraph.text.strip()
			
			if not text:
				continue
			
			# Check if it's a heading (based on style)
			is_heading = False
			heading_level = 1
			
			try:
				if paragraph.style and hasattr(paragraph.style, 'name') and paragraph.style.name:
					style_name = paragraph.style.name.lower()
					if 'heading' in style_name:
						is_heading = True
						try:
							# Extract heading level (e.g., "Heading 1" -> 1)
							level_str = style_name.replace('heading', '').strip()
							if level_str:
								heading_level = int(level_str)
						except (ValueError, AttributeError):
							heading_level = 1
			except (AttributeError, Exception):
				pass
			
			# Check if paragraph has bold formatting (likely heading)
			if not is_heading:
				try:
					for run in paragraph.runs:
						if hasattr(run, 'bold') and run.bold and len(text) < 100:  # Short bold text is likely a heading
							is_heading = True
							break
				except (AttributeError, Exception):
					pass
			
			if is_heading:
				# Save previous section if exists
				if current_section:
					sections.append(current_section)
				
				# Start new section
				section_id += 1
				# Store formatting info
				style_name = paragraph.style.name if paragraph.style and hasattr(paragraph.style, 'name') else None
				current_section = {
					'id': f'section-{section_id}',
					'title': text,
					'content': '',
					'type': 'heading',
					'order': section_id,
					'level': heading_level,
					'style': style_name,
					'formatting': _extract_paragraph_formatting(paragraph)
				}
			else:
				# Add to current section or create new one
				if not current_section:
					section_id += 1
					style_name = paragraph.style.name if paragraph.style and hasattr(paragraph.style, 'name') else None
					current_section = {
						'id': f'section-{section_id}',
						'title': f'Section {section_id}',
						'content': '',
						'type': 'paragraph',
						'order': section_id,
						'style': style_name,
						'formatting': _extract_paragraph_formatting(paragraph)
					}
				
				# Convert paragraph to HTML
				para_html = paragraph_to_html(paragraph)
				if current_section['content']:
					current_section['content'] += '<p>' + para_html + '</p>'
				else:
					current_section['content'] = '<p>' + para_html + '</p>'
		
		elif isinstance(element, CT_Tbl):
			table = Table(element, doc)
			
			# Save previous section if exists
			if current_section:
				sections.append(current_section)
				current_section = None
			
			# Create table section
			section_id += 1
			table_html = table_to_html(table)
			sections.append({
				'id': f'section-{section_id}',
				'title': f'Table {section_id}',
				'content': table_html,
				'type': 'table',
				'order': section_id
			})
	
	# Add last section
	if current_section:
		sections.append(current_section)
	
	return sections


def _extract_paragraph_formatting(paragraph: Paragraph) -> dict:
	"""Extract formatting information from paragraph"""
	formatting = {}
	try:
		if paragraph.style and hasattr(paragraph.style, 'name'):
			formatting['style'] = paragraph.style.name
		
		if paragraph.paragraph_format:
			if hasattr(paragraph.paragraph_format, 'alignment') and paragraph.paragraph_format.alignment:
				formatting['alignment'] = str(paragraph.paragraph_format.alignment)
		
		# Extract run formatting
		formatting['runs'] = []
		for run in paragraph.runs:
			run_format = {}
			if hasattr(run, 'font'):
				if hasattr(run.font, 'name') and run.font.name:
					run_format['font_name'] = run.font.name
				if hasattr(run.font, 'size') and run.font.size:
					run_format['font_size'] = str(run.font.size)
				if hasattr(run.font, 'color') and run.font.color and hasattr(run.font.color, 'rgb'):
					run_format['font_color'] = str(run.font.color.rgb)
				if hasattr(run.font, 'bold'):
					run_format['bold'] = run.font.bold
				if hasattr(run.font, 'italic'):
					run_format['italic'] = run.font.italic
				if hasattr(run.font, 'underline'):
					run_format['underline'] = run.font.underline
			run_format['text'] = run.text if hasattr(run, 'text') else ''
			formatting['runs'].append(run_format)
	except Exception as e:
		logger.warning(f"Error extracting formatting: {e}")
	
	return formatting


def paragraph_to_html(paragraph: Paragraph) -> str:
	"""Convert paragraph to HTML preserving ALL formatting (colors, fonts, sizes, etc.)"""
	from docx.shared import RGBColor
	import html
	
	html_parts = []
	
	try:
		for run in paragraph.runs:
			text = run.text if hasattr(run, 'text') else ''
			
			if not text:
				continue
			
			# Escape HTML entities
			text = html.escape(text)
			
			# Build style attribute with all formatting
			styles = []
			
			# Font properties
			if hasattr(run, 'font') and run.font:
				# Font name
				if hasattr(run.font, 'name') and run.font.name:
					styles.append(f"font-family: '{run.font.name}'")
				
				# Font size
				if hasattr(run.font, 'size') and run.font.size:
					# Convert from EMU to points (1 point = 12700 EMU)
					size_pt = run.font.size.pt if hasattr(run.font.size, 'pt') else run.font.size / 12700
					styles.append(f"font-size: {size_pt:.1f}pt")
				
				# Font color
				if hasattr(run.font, 'color') and run.font.color:
					if hasattr(run.font.color, 'rgb') and run.font.color.rgb:
						rgb = run.font.color.rgb
						if isinstance(rgb, RGBColor):
							# RGBColor stores as hex string like "FF0000"
							# Extract R, G, B values
							rgb_str = str(rgb)
							try:
								# Convert hex string to integer
								rgb_int = int(rgb_str, 16) if len(rgb_str) == 6 else int(rgb_str, 16) if rgb_str.startswith('0x') else int(rgb_str)
								r = (rgb_int >> 16) & 0xFF
								g = (rgb_int >> 8) & 0xFF
								b = rgb_int & 0xFF
								styles.append(f"color: rgb({r}, {g}, {b})")
							except (ValueError, AttributeError):
								# Fallback: try to parse as tuple if available
								pass
				
				# Background/highlight color
				if hasattr(run.font, 'highlight_color') and run.font.highlight_color:
					# python-docx has limited highlight support, but we can try
					pass
			
			# Text formatting
			formatting_tags = []
			if hasattr(run, 'bold') and run.bold:
				formatting_tags.append('strong')
			if hasattr(run, 'italic') and run.italic:
				formatting_tags.append('em')
			if hasattr(run, 'underline') and run.underline:
				formatting_tags.append('u')
			if hasattr(run, 'font') and run.font and hasattr(run.font, 'strike') and run.font.strike:
				formatting_tags.append('s')
			if hasattr(run, 'font') and run.font and hasattr(run.font, 'subscript') and run.font.subscript:
				formatting_tags.append('sub')
			if hasattr(run, 'font') and run.font and hasattr(run.font, 'superscript') and run.font.superscript:
				formatting_tags.append('sup')
			
			# Build HTML with style and formatting tags
			if styles:
				style_attr = '; '.join(styles)
				text = f'<span style="{style_attr}">{text}</span>'
			
			# Apply formatting tags (wrap from outside to inside)
			for tag in reversed(formatting_tags):
				text = f'<{tag}>{text}</{tag}>'
			
			html_parts.append(text)
			
	except (AttributeError, Exception) as e:
		logger.warning(f"Error processing paragraph runs: {e}")
		# Fallback to plain text with HTML escape
		if hasattr(paragraph, 'text'):
			return html.escape(paragraph.text)
		return ''
	
	result = ''.join(html_parts) if html_parts else ''
	
	# If paragraph has alignment, wrap in div with alignment
	if paragraph.paragraph_format and hasattr(paragraph.paragraph_format, 'alignment'):
		alignment = paragraph.paragraph_format.alignment
		if alignment:
			if 'RIGHT' in str(alignment):
				result = f'<div style="text-align: right;">{result}</div>'
			elif 'CENTER' in str(alignment):
				result = f'<div style="text-align: center;">{result}</div>'
			elif 'LEFT' in str(alignment):
				result = f'<div style="text-align: left;">{result}</div>'
	
	return result if result else (html.escape(paragraph.text) if hasattr(paragraph, 'text') else '')


def table_to_html(table: Table) -> str:
	"""Convert table to HTML preserving formatting"""
	from docx.shared import RGBColor
	import html
	
	html_parts = ['<table border="1" cellspacing="0" cellpadding="8" style="border-collapse: collapse; width: 100%; direction: rtl; text-align: right; margin: 10px 0;">']
	
	try:
		for row_idx, row in enumerate(table.rows):
			html_parts.append('<tr>')
			for cell_idx, cell in enumerate(row.cells):
				# Extract cell content with formatting
				cell_content = []
				for para in cell.paragraphs:
					para_html = paragraph_to_html(para)
					if para_html:
						# Remove outer div if paragraph_to_html wrapped it
						if para_html.startswith('<div'):
							# Extract inner content
							import re
							match = re.match(r'<div[^>]*>(.*?)</div>', para_html, re.DOTALL)
							if match:
								para_html = match.group(1)
						cell_content.append(para_html)
				
				cell_text = ' '.join(cell_content) if cell_content else html.escape(cell.text.strip() if cell.text else '')
				if not cell_text.strip():
					cell_text = '&nbsp;'
				
				# Get cell formatting
				cell_styles = ['direction: rtl;', 'text-align: right;', 'vertical-align: top;']
				if cell.paragraphs:
					first_para = cell.paragraphs[0]
					if first_para.paragraph_format and hasattr(first_para.paragraph_format, 'alignment'):
						alignment = first_para.paragraph_format.alignment
						if 'RIGHT' in str(alignment):
							cell_styles.append('text-align: right;')
						elif 'CENTER' in str(alignment):
							cell_styles.append('text-align: center;')
						elif 'LEFT' in str(alignment):
							cell_styles.append('text-align: left;')
				
				# Add border styles
				cell_styles.append('border: 1px solid #000;')
				cell_styles.append('padding: 8px;')
				
				style_attr = ' '.join(cell_styles)
				html_parts.append(f'<td style="{style_attr}">{cell_text}</td>')
			html_parts.append('</tr>')
	except Exception as e:
		logger.warning(f"Error processing table: {e}")
		# Fallback to simple table
		html_parts = ['<table border="1" cellspacing="0" cellpadding="8" style="border-collapse: collapse; width: 100%; direction: rtl; text-align: right; margin: 10px 0;">']
		for row in table.rows:
			html_parts.append('<tr>')
			for cell in row.cells:
				cell_text = html.escape(cell.text.strip() if cell.text else '')
				if not cell_text:
					cell_text = '&nbsp;'
				html_parts.append(f'<td style="direction: rtl; text-align: right; border: 1px solid #000; padding: 8px;">{cell_text}</td>')
			html_parts.append('</tr>')
	
	html_parts.append('</table>')
	return ''.join(html_parts)


@router.post("/parse-template")
async def parse_template(
	request: Request,
	template: UploadFile | None = File(default=None),
	file: UploadFile | None = File(default=None),
):
	"""Analyze a Word template and extract sections (same response as /word-template/analyze on main)."""
	if not _sent_by_reporting_node(request):
		logger.warning("parse-template: refused a call that did not come through reporting Node")
		return JSONResponse(status_code=403, content={'success': False, 'message': 'You do not have permission to perform this action.'})
	upload = template or file
	if upload is None:
		return _bad_request('No file provided')

	file_name = upload.filename or ''
	if not file_name.lower().endswith(('.docx', '.doc')):
		return _bad_request('Invalid file type. Expected a .docx file.')

	max_bytes = _mb_setting('UPLOAD_MAX_FILE_MB', 50)
	file_content = await upload.read(max_bytes + 1)
	if len(file_content) == 0:
		return _bad_request('File is empty')
	if len(file_content) > max_bytes:
		raise HTTPException(status_code=413, detail='File too large')

	# A .docx is a zip: refuse ones that unpack to far more than they weigh before python-docx opens them.
	try:
		with zipfile.ZipFile(BytesIO(file_content)) as zf:
			unzipped = sum(info.file_size for info in zf.infolist())
	except zipfile.BadZipFile:
		return _bad_request('Failed to open Word document. Please ensure the file is a valid .docx file.')
	if unzipped > _mb_setting('WORD_TEMPLATE_MAX_UNZIPPED_MB', 200):
		raise HTTPException(status_code=413, detail='File too large')

	try:
		doc = Document(BytesIO(file_content))
	except Exception as doc_error:
		logger.error("parse-template: cannot open %s: %s", file_name, doc_error)
		return _bad_request('Failed to open Word document. Please ensure the file is a valid .docx file.')

	try:
		sections = extract_sections_from_docx(doc)
	except Exception:
		logger.exception("parse-template: section extraction failed for %s", file_name)
		return JSONResponse(status_code=500, content={'success': False, 'error': 'Failed to extract sections from the template.'})

	logger.info("Extracted %d sections from template", len(sections))

	# Ensure at least one section exists
	if not sections:
		sections = [{
			'id': 'section-1',
			'title': 'Document Content',
			'content': '<p>No structured sections found. Document content will be available for editing.</p>',
			'type': 'paragraph',
			'order': 1
		}]

	all_placeholders, total_placeholders = extract_all_placeholders_from_docx_bytes(file_content)
	return JSONResponse(
		status_code=200,
		content={
			'success': True,
			'sections': sections,
			'total_sections': len(sections),
			'all_placeholders': all_placeholders,
			'total_placeholders': total_placeholders,
		}
	)
